"""
divisor.py - Fase 1: División y Almacenamiento (El Splitter)

Abre un documento Word Maestro de especificaciones técnicas, identifica los títulos
de cada partida mediante expresiones regulares basadas en numeración (ej. 01.01.03 o 1.2.4),
aísla el bloque completo (párrafos, tablas, imágenes, estilos) sin perder formato,
lo guarda como archivo independiente en /partidas_db/ y registra los metadatos en SQLite.

Soporta dos motores:
1. 'docx' (Nativo Multiplataforma - Linux / Windows / Mac):
   Clona la estructura OpenXML del documento maestro, conservando estilos, numeraciones,
   tablas y medios (imágenes/gráficos) al 100% mediante poda quirúrgica del cuerpo XML.
2. 'win32com' (Nativo Windows con Microsoft Word instalado):
   Manipula la aplicación Word por interfaz COM, copiando y pegando rangos nativos.
"""

import os
import sys
import re
import shutil
import argparse
from typing import List, Dict, Optional, Tuple, Any

import docx
from db_manager import init_db, guardar_partidas_batch, listar_todas_las_partidas

# Expresión regular por defecto: Detecta numeración como 01.01.03, 1.2.4, 02.01, etc.
# y captura el ID numérico en grupo 1 y el nombre/descripción en grupo 2.
DEFAULT_REGEX = r'^\s*(\d+(?:\.\d+)+)\s*[-.:]?\s*(.*)$'


def detectar_backend_disponible(preferencia: str = "auto") -> str:
    """Determina si se usa win32com (Windows) o docx (Multiplataforma)."""
    if preferencia == "win32com":
        try:
            import win32com.client
            return "win32com"
        except ImportError:
            print("[ADVERTENCIA] win32com no está instalado. Se utilizará el motor 'docx'.")
            return "docx"
            
    if preferencia == "docx":
        return "docx"
        
    # Modo 'auto': Si está en Windows y win32com está disponible, usa win32com; si no, docx.
    if sys.platform.startswith("win"):
        try:
            import win32com.client
            return "win32com"
        except ImportError:
            pass
            
    return "docx"


def contar_elementos_en_docx(doc_path: str) -> Dict[str, int]:
    """Inspecciona un archivo docx y retorna conteo de párrafos, tablas e imágenes."""
    try:
        doc = docx.Document(doc_path)
        num_parrafos = len(doc.paragraphs)
        num_tablas = len(doc.tables)
        num_imagenes = len(doc.inline_shapes)
        return {"parrafos": num_parrafos, "tablas": num_tablas, "imagenes": num_imagenes}
    except Exception:
        return {"parrafos": 0, "tablas": 0, "imagenes": 0}


def dividir_con_docx(
    master_path: str,
    output_dir: str = "partidas_db",
    regex_pattern: str = DEFAULT_REGEX
) -> List[Dict[str, Any]]:
    """
    Divide el documento maestro utilizando manipulación de OpenXML nativo.
    
    Estrategia de conservación total de formato:
    1. Abre el documento maestro para inspeccionar el cuerpo XML (w:body).
    2. Identifica los elementos de párrafo que coinciden con la expresión regular de partidas.
    3. Para cada partida, copia el archivo maestro completo como base (preservando estilos,
       fuentes, viñetas en numbering.xml, imágenes en word/media/ y relaciones rId).
    4. Poda los elementos XML externos al bloque de la partida, reteniendo las propiedades
       de sección (sectPr) al final para mantener márgenes y geometría de página.
    """
    os.makedirs(output_dir, exist_ok=True)
    doc_maestro = docx.Document(master_path)
    body = doc_maestro._body._element
    children = list(body)
    
    comp_regex = re.compile(regex_pattern, re.IGNORECASE)
    partidas_detectadas = []
    
    # 1. Identificar posiciones de los encabezados de partida
    for idx, child in enumerate(children):
        if child.tag.endswith('p'):
            # Convertir elemento a párrafo para leer texto completo
            p_obj = docx.text.paragraph.Paragraph(child, doc_maestro)
            texto = p_obj.text.strip()
            match = comp_regex.match(texto)
            if match:
                id_partida = match.group(1).strip()
                nombre_partida = match.group(2).strip()
                # Si el título no tiene texto adicional en la misma línea, usar el ID o buscar siguiente línea
                if not nombre_partida:
                    nombre_partida = f"PARTIDA {id_partida}"
                partidas_detectadas.append({
                    "child_idx": idx,
                    "id_partida": id_partida,
                    "nombre_partida": nombre_partida,
                    "texto_original": texto
                })
                
    if not partidas_detectadas:
        print(f"[ERROR] No se encontraron partidas con el patrón '{regex_pattern}' en '{master_path}'.")
        return []
        
    print(f"[INFO] Se identificaron {len(partidas_detectadas)} partidas en el documento maestro.")
    
    # Verificar si el último elemento es sectPr
    has_sectPr = children[-1].tag.endswith('sectPr')
    last_body_idx = len(children) - 1 if has_sectPr else len(children)
    
    resultados = []
    
    # 2. Extraer cada bloque quirúrgicamente
    for i, partida in enumerate(partidas_detectadas):
        id_partida = partida["id_partida"]
        nombre = partida["nombre_partida"]
        start_idx = partida["child_idx"]
        
        # El bloque termina donde empieza la siguiente partida o al final del contenido
        if i + 1 < len(partidas_detectadas):
            end_idx = partidas_detectadas[i + 1]["child_idx"]
        else:
            end_idx = last_body_idx
            
        # Conjunto de índices a conservar (incluyendo sectPr final si existe)
        keep_indices = set(range(start_idx, end_idx))
        if has_sectPr:
            keep_indices.add(len(children) - 1)
            
        # Nombre de archivo estandarizado con guiones bajos (ej. 01_01_03.docx)
        filename = f"{id_partida.replace('.', '_')}.docx"
        out_path = os.path.join(output_dir, filename)
        
        # Clonar el archivo maestro para conservar todos los paquetes, estilos e imágenes
        shutil.copyfile(master_path, out_path)
        
        # Abrir el clon y podar los elementos no pertenecientes a la partida
        part_doc = docx.Document(out_path)
        part_body = part_doc._body._element
        part_children = list(part_body)
        
        for c_idx, child in enumerate(part_children):
            if c_idx not in keep_indices:
                part_body.remove(child)
                
        part_doc.save(out_path)
        
        # Estadísticas del documento generado
        stats = contar_elementos_en_docx(out_path)
        
        resultados.append({
            "id_partida": id_partida,
            "nombre_partida": nombre,
            "ruta_archivo": os.path.normpath(out_path),
            "stats": stats
        })
        
        print(f"  ✓ [{id_partida}] {nombre[:45]:<45} -> {filename} "
              f"({stats['parrafos']} párr, {stats['tablas']} tablas, {stats['imagenes']} imgs)")
        
    return resultados


def dividir_con_win32com(
    master_path: str,
    output_dir: str = "partidas_db",
    regex_pattern: str = DEFAULT_REGEX
) -> List[Dict[str, Any]]:
    """
    Divide el documento maestro utilizando la interfaz nativa COM de Windows (Word.Application).
    """
    try:
        import win32com.client
        import pythoncom
    except ImportError:
        raise RuntimeError("win32com no está instalado. Instale 'pywin32' o use backend 'docx'.")

    os.makedirs(output_dir, exist_ok=True)
    abs_master = os.path.abspath(master_path)
    abs_out_dir = os.path.abspath(output_dir)
    comp_regex = re.compile(regex_pattern, re.IGNORECASE)

    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = False
    
    resultados = []
    try:
        doc = word.Documents.Open(abs_master, ReadOnly=True)
        partidas_info = []
        
        print("[INFO] Analizando párrafos con Microsoft Word COM...")
        total_paras = doc.Paragraphs.Count
        for i in range(1, total_paras + 1):
            p = doc.Paragraphs(i)
            txt = p.Range.Text.strip()
            match = comp_regex.match(txt)
            if match:
                id_partida = match.group(1).strip()
                nombre_partida = match.group(2).strip() or f"PARTIDA {id_partida}"
                partidas_info.append({
                    "id_partida": id_partida,
                    "nombre_partida": nombre_partida,
                    "start_pos": p.Range.Start,
                    "para_idx": i
                })
                
        if not partidas_info:
            print(f"[ERROR] No se encontraron partidas con el patrón '{regex_pattern}' en Word COM.")
            doc.Close(False)
            return []
            
        print(f"[INFO] Se detectaron {len(partidas_info)} partidas. Extrayendo rangos nativos...")
        
        for i, partida in enumerate(partidas_info):
            start_pos = partida["start_pos"]
            if i + 1 < len(partidas_info):
                end_pos = partidas_info[i + 1]["start_pos"]
            else:
                end_pos = doc.Content.End
                
            rng = doc.Range(Start=start_pos, End=end_pos)
            rng.Copy()
            
            new_doc = word.Documents.Add()
            new_doc.Content.Paste()
            
            id_partida = partida["id_partida"]
            filename = f"{id_partida.replace('.', '_')}.docx"
            out_file = os.path.join(abs_out_dir, filename)
            
            # Formato 16 = wdFormatDocumentDefault (.docx)
            new_doc.SaveAs(out_file, FileFormat=16)
            new_doc.Close(False)
            
            rel_path = os.path.relpath(out_file)
            stats = contar_elementos_en_docx(out_file)
            resultados.append({
                "id_partida": id_partida,
                "nombre_partida": partida["nombre_partida"],
                "ruta_archivo": os.path.normpath(rel_path),
                "stats": stats
            })
            print(f"  ✓ [COM] [{id_partida}] {partida['nombre_partida'][:40]} -> {filename}")
            
        doc.Close(False)
    finally:
        word.Quit()
        pythoncom.CoUninitialize()
        
    return resultados


def ejecutar_division(
    master_path: str,
    db_path: str = "partidas.db",
    output_dir: str = "partidas_db",
    regex_pattern: str = DEFAULT_REGEX,
    backend: str = "auto",
    reset_db: bool = False
) -> List[Dict[str, Any]]:
    """
    Función principal para la Fase 1.
    Abre el documento maestro, extrae las partidas y las almacena en disco y base de datos.
    """
    if not os.path.isfile(master_path):
        raise FileNotFoundError(f"El documento maestro no existe: '{master_path}'")
        
    motor = detectar_backend_disponible(backend)
    print(f"==================================================")
    print(f" FASE 1: DIVISIÓN Y ALMACENAMIENTO (EL SPLITTER)")
    print(f" Documento Maestro : {master_path}")
    print(f" Carpeta Destino   : {output_dir}")
    print(f" Base de Datos     : {db_path}")
    print(f" Motor Seleccionado: {motor.upper()}")
    print(f" Patrón Regex      : {regex_pattern}")
    print(f"==================================================")
    
    if motor == "win32com":
        partidas = dividir_con_win32com(master_path, output_dir, regex_pattern)
    else:
        partidas = dividir_con_docx(master_path, output_dir, regex_pattern)
        
    if not partidas:
        print("[AVISO] No se generaron partidas.")
        return []
        
    # Inicializar e indexar en la base de datos SQLite
    init_db(db_path, reset=reset_db)
    total_guardadas = guardar_partidas_batch(db_path, partidas)
    
    print("--------------------------------------------------")
    print(f"✔ Proceso completado exitosamente:")
    print(f"  Total de partidas extraídas : {len(partidas)}")
    print(f"  Total indexadas en SQLite   : {total_guardadas}")
    print(f"  Ubicación de archivos .docx : {os.path.abspath(output_dir)}")
    print(f"  Base de datos actualizada   : {os.path.abspath(db_path)}")
    print("--------------------------------------------------")
    
    return partidas


def main():
    parser = argparse.ArgumentParser(
        description="Fase 1: Divide un documento maestro de especificaciones en archivos individuales."
    )
    parser.add_argument(
        "-m", "--maestro", required=True,
        help="Ruta al archivo Word Maestro (.docx)."
    )
    parser.add_argument(
        "-o", "--output-dir", default="partidas_db",
        help="Directorio donde se guardarán los archivos .docx de cada partida (default: partidas_db)."
    )
    parser.add_argument(
        "-d", "--db", default="partidas.db",
        help="Ruta a la base de datos SQLite (default: partidas.db)."
    )
    parser.add_argument(
        "-r", "--regex", default=DEFAULT_REGEX,
        help="Expresión regular para identificar los encabezados de partida."
    )
    parser.add_argument(
        "-b", "--backend", choices=["auto", "docx", "win32com"], default="auto",
        help="Motor de división: 'auto' (detecta SO), 'docx' (multiplataforma) o 'win32com' (Windows COM)."
    )
    parser.add_argument(
        "--reset-db", action="store_true",
        help="Reinicia la tabla de partidas en la base de datos antes de registrar."
    )
    
    args = parser.parse_args()
    
    try:
        ejecutar_division(
            master_path=args.maestro,
            db_path=args.db,
            output_dir=args.output_dir,
            regex_pattern=args.regex,
            backend=args.backend,
            reset_db=args.reset_db
        )
    except Exception as e:
        print(f"\n[ERROR CRÍTICO] {e}", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
